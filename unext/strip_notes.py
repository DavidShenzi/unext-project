"""Produce a speaker-note-free copy of a deck.

The presentation carries its speaker notes in the notes pane of every slide -- roughly
33,000 characters of them. Those notes are working material: they contain the phrasing
we rehearse with, the answers we plan to give under questioning, and candid statements
about our own mistakes. None of that belongs in a public repository or in a file handed
to anyone but us.

python-pptx has no "delete the notes" call. Emptying `notes_text_frame.text` leaves the
notesSlide part in the package, so the relationship and an empty frame survive; that is
enough for the text to be gone, but it leaves the structure behind. This removes the
notesSlide part itself and drops the relationship from the slide, so a reader opening
the file sees no notes pane content and the XML carries nothing to recover.

Verification is part of the job, not an afterthought: the output is reopened and every
slide checked, and the script fails loudly rather than silently producing a file that
still carries notes.

Usage:
    python strip_notes.py in.pptx out.pptx
"""

import shutil
import sys

from pptx import Presentation


def strip(src, dst):
    """Copy `src` to `dst` with every notes slide removed. Returns slides stripped."""
    shutil.copyfile(src, dst)
    prs = Presentation(dst)
    n = 0
    for slide in prs.slides:
        if not slide.has_notes_slide:
            continue
        # Drop the notes slide part and the relationship pointing at it. Clearing the
        # text alone would leave the part in the package.
        rels = slide.part.rels
        drop = [rid for rid, rel in rels.items()
                if rel.reltype.endswith('/notesSlide')]
        for rid in drop:
            rels._rels.pop(rid, None)
            n += 1
    prs.save(dst)
    return n


def verify(path):
    """Reopen and confirm nothing carries notes text. Returns list of offending slides."""
    prs = Presentation(path)
    bad = []
    for i, slide in enumerate(prs.slides, 1):
        if slide.has_notes_slide:
            t = slide.notes_slide.notes_text_frame.text.strip()
            if t:
                bad.append((i, len(t)))
    return bad


def main():
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    src, dst = sys.argv[1], sys.argv[2]
    n = strip(src, dst)
    bad = verify(dst)
    if bad:
        raise SystemExit(f'FAILED: {len(bad)} slide(s) still carry notes: {bad}')
    print(f'{dst}: stripped {n} notes slide(s), verified clean')


if __name__ == '__main__':
    main()
